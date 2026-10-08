from pathlib import Path
from docx import Document
from docx.shared import RGBColor, Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table

path = Path(__file__).parent / '_output/infosec-intro-lab05-report.docx'
doc = Document(path)
for style in doc.styles:
    if style.type == 1 or style.type == 2:
        style.font.color.rgb = RGBColor(0, 0, 0)
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if paragraph.style.name in ('Title', 'Subtitle') or paragraph.style.name.startswith('Heading'):
            run.font.color.rgb = RGBColor(0, 0, 0)
    if paragraph._p.xpath('.//w:drawing'):
        paragraph.paragraph_format.keep_with_next = True
    if paragraph.style.name in ('Caption', 'Image Caption'):
        paragraph.paragraph_format.keep_with_next = False
        paragraph.paragraph_format.keep_together = True
    if paragraph.style.name == 'Bibliography':
        paragraph.paragraph_format.keep_together = True
for paragraph in doc.paragraphs:
    if paragraph.text == 'Список литературы':
        paragraph.paragraph_format.page_break_before = True
for table_element in doc._element.body.xpath('.//w:tbl'):
    table = Table(table_element, doc)
    if table._tbl.xpath('.//w:drawing') or table._tbl.xpath('./w:tr/w:tc/w:tbl'):
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for i, p in enumerate(cell.paragraphs):
                    p.paragraph_format.keep_with_next = i < len(cell.paragraphs)-1
                    p.paragraph_format.keep_together = True
        continue
    table.autofit = False
    columns = len(table.columns)
    widths = [0.65, 1.25, 1.15, 1.15, 1.15, 1.15] if columns == 6 else [3.0, 1.75, 1.75] if columns == 3 else []
    if len(widths) == columns:
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    properties = table._tbl.tblPr
    for tag in ('tblBorders', 'tblCellMar'):
        current = properties.find(qn('w:' + tag))
        if current is not None:
            properties.remove(current)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        item = OxmlElement('w:' + edge)
        item.set(qn('w:val'), 'single')
        item.set(qn('w:sz'), '4')
        item.set(qn('w:color'), 'D9D9D9')
        borders.append(item)
    properties.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge, value in [('top', 80), ('bottom', 80), ('left', 80), ('right', 80)]:
        item = OxmlElement('w:' + edge)
        item.set(qn('w:w'), str(value))
        item.set(qn('w:type'), 'dxa')
        margins.append(item)
    properties.append(margins)
    for r, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if r == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for c, cell in enumerate(row.cells):
            cell.vertical_alignment = 1
            if r == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8EDF2')
                cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.keep_with_next = r < len(table.rows)-1
                for run in p.runs:
                    run.font.size = Pt(9 if columns > 5 else 10)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.bold = r == 0
doc.save(path)
